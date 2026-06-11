import os
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_diagrams():
    os.makedirs('ppt_assets', exist_ok=True)
    
    # 1. Data Distribution Pie Chart
    plt.figure(figsize=(6, 5), facecolor='white')
    labels = ['Perfect RIS\n(8,000 samples)', 'Faulty RIS\n(8,000 samples)']
    sizes = [8000, 8000]
    colors = ['#2ecc71', '#e74c3c']
    plt.pie(sizes, labels=labels, colors=colors, autopct='%1.1f%%', startangle=90, textprops={'fontsize': 14, 'weight': 'bold'})
    plt.title('Total Dataset: 16,000 Samples', fontsize=16, weight='bold')
    plt.savefig('ppt_assets/data_dist.png', bbox_inches='tight', dpi=300)
    plt.close()

    # 2. Workflow Diagram
    plt.figure(figsize=(9, 3), facecolor='white')
    plt.axis('off')
    boxes = ['1. Virtual\nSimulation', '2. Signal\nTranslation', '3. AI\nTraining', '4. Fault\nDetection', '5. Fault\nLocalization']
    colors_box = ['#3498db', '#9b59b6', '#f1c40f', '#e67e22', '#e74c3c']
    for i, box in enumerate(boxes):
        plt.text(i*2, 0.5, box, ha='center', va='center', fontsize=12, weight='bold', color='white' if i!=2 else 'black',
                 bbox=dict(boxstyle="round,pad=0.5", facecolor=colors_box[i], edgecolor="black", lw=2, alpha=0.9))
        if i < len(boxes)-1:
            plt.arrow(i*2 + 0.8, 0.5, 0.4, 0, head_width=0.1, head_length=0.1, fc='k', ec='k')
    plt.xlim(-1, 9)
    plt.ylim(0, 1)
    plt.savefig('ppt_assets/workflow.png', bbox_inches='tight', dpi=300)
    plt.close()

    # 3. Tasks Diagram
    plt.figure(figsize=(7, 4), facecolor='white')
    plt.axis('off')
    plt.text(0.5, 0.8, 'Received Wireless Signal', ha='center', va='center', fontsize=14, weight='bold', bbox=dict(boxstyle="round,pad=0.5", facecolor="#ecf0f1", edgecolor="black", lw=2))
    plt.arrow(0.5, 0.7, -0.2, -0.2, head_width=0.05, fc='k', ec='k')
    plt.arrow(0.5, 0.7, 0.2, -0.2, head_width=0.05, fc='k', ec='k')
    plt.text(0.3, 0.4, 'Task 1: Detection\n(Is it Broken?)', ha='center', va='center', fontsize=12, weight='bold', color='white', bbox=dict(boxstyle="round,pad=0.5", facecolor="#e74c3c", edgecolor="black", lw=2))
    plt.text(0.7, 0.4, 'Task 2: Localization\n(Which Tile?)', ha='center', va='center', fontsize=12, weight='bold', color='white', bbox=dict(boxstyle="round,pad=0.5", facecolor="#2980b9", edgecolor="black", lw=2))
    plt.xlim(0, 1)
    plt.ylim(0, 1)
    plt.savefig('ppt_assets/tasks.png', bbox_inches='tight', dpi=300)
    plt.close()

def add_slide(prs, title_text, content_bullets=None, image_path=None, image_top=3.0, image_left=1.5, image_width=7.0):
    slide_layout = prs.slide_layouts[1] if content_bullets else prs.slide_layouts[5] # 5 is title only
    slide = prs.slides.add_slide(slide_layout)
    
    # Apply a modern dark theme background (Dark Slate Blue)
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = RGBColor(30, 40, 60)

    # Title formatting
    title_shape = slide.shapes.title
    title_shape.text = title_text
    title_tf = title_shape.text_frame
    for p in title_tf.paragraphs:
        p.font.color.rgb = RGBColor(241, 196, 15) # Gold/Yellow for titles
        p.font.bold = True
        p.font.size = Pt(40)

    # Content bullets formatting
    if content_bullets and len(slide.placeholders) > 1:
        body_shape = slide.placeholders[1]
        tf = body_shape.text_frame
        tf.clear()
        for bullet in content_bullets:
            p = tf.add_paragraph()
            p.text = bullet
            p.font.color.rgb = RGBColor(236, 240, 241) # Off-white text
            p.font.size = Pt(24)
            p.space_after = Pt(14)

    # Add image if provided
    if image_path and os.path.exists(image_path):
        slide.shapes.add_picture(image_path, Inches(image_left), Inches(image_top), width=Inches(image_width))
        
    return slide

def create_presentation():
    create_diagrams()
    prs = Presentation()
    
    # --- Slide 1: Title Slide ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(20, 25, 40)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AI for Faulty RIS Element Prediction"
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(241, 196, 15)
    title.text_frame.paragraphs[0].font.bold = True
    subtitle.text = "A Deep Dive into Classical & Quantum Machine Learning\n\nCode Walkthrough & Project Defense"
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(200, 210, 220)

    # --- Slide 2: Project Overview ---
    add_slide(prs, "1. Project Overview", [
        "Modern 6G networks use Reconfigurable Intelligent Surfaces (RIS).",
        "Think of RIS as a massive 'Smart Mirror' made of tiny tiles.",
        "These mirrors bounce wireless signals perfectly to your phone.",
        "Goal: Use AI to automatically detect if a tile breaks."
    ])

    # --- Slide 3: The Core Problem ---
    add_slide(prs, "2. The Core Problem", [
        "RIS devices sit outdoors and face harsh weather.",
        "Individual reflection tiles can and will break down.",
        "When tiles break, the signal bounces wrong, and internet speed drops.",
        "Manual inspection of thousands of tiny tiles on a building is impossible."
    ])

    # --- Slide 4: Our Solution ---
    add_slide(prs, "3. Our Solution Approach", [
        "We can't visually inspect the mirror, but we can 'listen' to it.",
        "We analyze the wireless signal received by the user.",
        "We train Artificial Intelligence to recognize the 'sound' of a broken mirror.",
        "The AI will tell us EXACTLY which tile is broken just from the signal."
    ])

    # --- Slide 5: Workflow Architecture ---
    add_slide(prs, "4. Project Workflow Overview", [], 'ppt_assets/workflow.png', image_top=2.5, image_left=0.5, image_width=9)

    # --- Slide 6: Step 1 - The Simulation ---
    add_slide(prs, "5. Step 1: The Virtual World", [
        "We don't have a physical broken smart mirror.",
        "Instead, we wrote code to mathematically simulate the physical world.",
        "The code simulates antennas, the RIS, the user, and signal bouncing."
    ])

    # --- Slide 7: Code Breakdown - The Simulator ---
    add_slide(prs, "6. Code Breakdown: generate_data.py", [
        "What it does: Acts as our 'Video Game Engine' for physics.",
        "Key Code Variables:",
        " > N_Tx = 4 : Number of antennas sending the signal.",
        " > N_RIS = 8 : Elements on the mirror.",
        " > frequency = 26e9 : The pitch of our wireless signal."
    ])

    # --- Slide 8: Injecting the Faults ---
    add_slide(prs, "7. Injecting the Hardware Faults", [
        "How do we teach the AI what a broken mirror looks like?",
        "We intentionally 'break' the mirror in our simulation.",
        "Code Analogy: We literally turn off the reflection for one specific tile."
    ])

    # --- Slide 9: Code Breakdown - Breaking the Mirror ---
    add_slide(prs, "8. Code Breakdown: The Fault Simulator", [
        "Look at this line in the code:",
        " > phi1[fault_idx] = 0.0",
        "What it means:",
        " > 'phi1' is the array of mirror tiles.",
        " > 'fault_idx' is the exact tile we are breaking.",
        " > '= 0.0' means we turn its reflection power to zero (broken!)."
    ])

    # --- Slide 10: The Dataset Generated ---
    add_slide(prs, "9. The Dataset Size", [
        "The AI needs massive amounts of examples to learn.",
        "Total Dataset: 16,000 signal examples.",
        "We balanced the data so the AI doesn't get biased."
    ], 'ppt_assets/data_dist.png', image_top=3.5, image_left=2.5, image_width=5)

    # --- Slide 11: Data Preprocessing ---
    add_slide(prs, "10. Step 2: Data Preprocessing", [
        "Wireless signals are recorded as 'Complex Numbers' (Real + Imaginary).",
        "Standard AI models are not very good at reading complex numbers.",
        "We must translate this data into a language the AI understands."
    ])

    # --- Slide 12: Code Breakdown - Preprocessing ---
    add_slide(prs, "11. Code Breakdown: Signal Translation", [
        "Look at the 'preprocess(complex_data)' function:",
        " > np.real(complex_data)",
        " > np.imag(complex_data)",
        "What it does:",
        " > It splits the signal into two plain numbers.",
        " > Analogy: Like translating French into English so the AI can read it."
    ])

    # --- Slide 13: Step 3 - The AI Models ---
    add_slide(prs, "12. Step 3: The AI Brains", [
        "We tested three different AI 'Brains' to see which is smartest:",
        "1. MLP (Multi-Layer Perceptron) - The Standard Brain",
        "2. CNN (Convolutional Neural Network) - The Pattern Finder",
        "3. QML (Quantum Machine Learning) - The Futuristic Brain"
    ])

    # --- Slide 14: Model 1 - MLP ---
    add_slide(prs, "13. AI Model 1: MLP", [
        "Multi-Layer Perceptron (MLP).",
        "This is a classic, standard Neural Network.",
        "Pros: Simple, fast to train, well-understood.",
        "Cons: Might struggle if the signal patterns are extremely complex."
    ])

    # --- Slide 15: Model 2 - CNN ---
    add_slide(prs, "14. AI Model 2: CNN", [
        "Convolutional Neural Network (CNN).",
        "Usually used for Facial Recognition and Images.",
        "Why use it here?",
        " > We treat the signal like an 'image' to find hidden spatial patterns.",
        " > CNNs are excellent at spotting tiny anomalies."
    ])

    # --- Slide 16: Model 3 - QML ---
    add_slide(prs, "15. AI Model 3: QML", [
        "Quantum Machine Learning (QML).",
        "A hybrid model using Quantum Computing concepts.",
        "Why use it here?",
        " > Quantum models can process complex possibilities simultaneously.",
        " > It represents cutting-edge research in telecommunications."
    ])

    # --- Slide 17: Step 4 - The Two Tasks ---
    add_slide(prs, "16. Step 4: The Two AI Tasks", [], 'ppt_assets/tasks.png', image_top=2.5, image_left=1.5, image_width=7)

    # --- Slide 18: Training the AI ---
    add_slide(prs, "17. Training the Models (train.py)", [
        "We send our 16,000 samples into the AI models.",
        "We hide the answers, let the AI guess, and then grade its test.",
        "Over time (Epochs), the AI adjusts its brain to guess better."
    ])

    # --- Slide 19: Code Breakdown - The Teacher ---
    add_slide(prs, "18. Code Breakdown: The Loss Function", [
        "Look at this code in train.py:",
        " > criterion = nn.CrossEntropyLoss()",
        "What it does:",
        " > This is the 'Teacher' grading the AI's test.",
        " > If the AI guesses wrong, the 'Loss' goes up.",
        " > The AI's only goal is to make the Loss as close to 0 as possible."
    ])

    # --- Slide 20: Code Breakdown - The Optimizer ---
    add_slide(prs, "19. Code Breakdown: The Optimizer", [
        "Look at this code in train.py:",
        " > optimizer = torch.optim.Adam(model.parameters())",
        "What it does:",
        " > This is the mechanism that actually 'rewires' the AI's brain.",
        " > After the Teacher (Loss) tells the AI it failed, Adam (Optimizer)",
        "   adjusts the math so the AI does better on the next try."
    ])

    # --- Slide 21: Evaluation Metrics ---
    add_slide(prs, "20. Evaluation Metrics", [
        "How do we know if the AI is actually smart?",
        "1. Accuracy: Out of 100 tests, how many did it get perfectly right?",
        "2. F1-Score: A balanced score that ensures the AI isn't just guessing 'Healthy' every time.",
        "The code prints a final scorecard comparing MLP, CNN, and QML."
    ])

    # --- Slide 22: Conclusion ---
    add_slide(prs, "21. Conclusion", [
        "We successfully simulated a 6G smart mirror environment.",
        "We proved that we can find hardware failures without physical inspection.",
        "We compared classical AI against futuristic Quantum AI.",
        "This project paves the way for self-healing wireless networks."
    ])

    # --- Slide 23: Q & A ---
    add_slide(prs, "22. Questions & Answers", [
        "Thank you for your time.",
        "",
        "Open for Questions regarding:",
        " > The Simulation Math",
        " > The Neural Network Architectures",
        " > The Quantum Machine Learning approach"
    ])

    prs.save('Detailed_Stylish_Presentation.pptx')
    print("Presentation saved successfully as 'Detailed_Stylish_Presentation.pptx'")

if __name__ == '__main__':
    create_presentation()
