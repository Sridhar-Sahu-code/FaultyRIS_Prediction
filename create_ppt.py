from pptx import Presentation
from pptx.util import Inches, Pt
import os

def create_presentation():
    prs = Presentation()
    
    # Title Slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Faulty RIS Element Prediction"
    subtitle.text = "Using Classical and Quantum Machine Learning\n\nPrepared for Supervisor Review"
    
    # Slide 1: Introduction to RIS
    bullet_slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    title.text = "1. Introduction to RIS"
    tf = body.text_frame
    tf.text = "Reconfigurable Intelligent Surfaces (RIS) are 'smart mirrors' for wireless signals."
    p = tf.add_paragraph()
    p.text = "They consist of numerous small reflecting elements."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "They can reflect signals in specific directions to improve coverage and signal strength in 6G networks."
    p.level = 1

    # Slide 2: The Problem
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    title.text = "2. The Problem Statement"
    tf = body.text_frame
    tf.text = "Individual elements on the RIS can break down or become faulty."
    p = tf.add_paragraph()
    p.text = "A faulty element reflects signals improperly, degrading overall network performance."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Manually checking elements is impossible in real-time."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Goal: Use Artificial Intelligence to automatically detect and locate these faulty elements."
    p.level = 1

    # Slide 3: Dataset Generation
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    title.text = "3. Dataset Generation"
    tf = body.text_frame
    tf.text = "We mathematically simulate the wireless signals between the transmitter, RIS, and receiver."
    p = tf.add_paragraph()
    p.text = "Total Dataset Size: 16,000 samples."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "8,000 'Perfect' samples: The RIS is fully functional."
    p.level = 2
    p = tf.add_paragraph()
    p.text = "8,000 'Faulty' samples: One of the 16 elements is broken (500 samples per fault location)."
    p.level = 2

    # Slide 4: AI Models Used
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    title.text = "4. AI Models Used"
    tf = body.text_frame
    tf.text = "We compare three distinct Artificial Intelligence architectures:"
    p = tf.add_paragraph()
    p.text = "MLP (Multi-Layer Perceptron): Standard Neural Network baseline."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "CNN (Convolutional Neural Network): Better at identifying spatial patterns in the signal data."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "QML (Quantum Machine Learning): A hybrid quantum-classical model that exploits quantum computing concepts."
    p.level = 1

    # Slide 5: The Two Tasks
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    title.text = "5. The Two Main Tasks"
    tf = body.text_frame
    tf.text = "Task 1: Fault Detection (Binary Classification)"
    p = tf.add_paragraph()
    p.text = "Is the RIS faulty or perfect? (Yes/No)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Task 2: Fault Localization (16-Class Classification)"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "If it is faulty, which exact element (out of 16) is broken?"
    p.level = 1

    # Slide 6: Conclusion
    slide = prs.slides.add_slide(bullet_slide_layout)
    title = slide.shapes.title
    body = slide.placeholders[1]
    title.text = "6. Conclusion"
    tf = body.text_frame
    tf.text = "By using deep learning and quantum machine learning:"
    p = tf.add_paragraph()
    p.text = "We can accurately detect the presence of a fault."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "We can pinpoint the exact location of the hardware failure."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "This allows for automated maintenance and self-healing smart wireless networks."
    p.level = 1

    prs.save('RIS_Fault_Prediction_Presentation.pptx')
    print("Presentation saved successfully as 'RIS_Fault_Prediction_Presentation.pptx'")

if __name__ == '__main__':
    create_presentation()
